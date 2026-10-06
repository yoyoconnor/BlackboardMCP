"""End-to-end test of every tool against a fake Blackboard API (no network)."""

import asyncio
import io
import os
import sys
import tempfile
from datetime import datetime, timedelta, timezone

import httpx

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ["BB_CACHE_DIR"] = tempfile.mkdtemp()

from blackboard_mcp import server  # noqa: E402
from blackboard_mcp.client import BlackboardClient  # noqa: E402

NOW = datetime.now(timezone.utc)
iso = lambda d: d.strftime("%Y-%m-%dT%H:%M:%S.000Z")  # noqa: E731


def make_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for i, (t, b, n) in enumerate([
        ("TCP Congestion Control", "Slow start doubles cwnd each RTT", "Remember AIMD"),
        ("Fast Retransmit", "Triggered by 3 duplicate ACKs", ""),
    ]):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text_frame.text = b
        if n:
            s.notes_slide.notes_text_frame.text = n
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(612, 792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


PPTX, PDF = make_pptx(), make_pdf()
API = "/learn/api/public"
ROUTES = {
    f"{API}/v1/users/me": {"id": "_9_1", "userName": "student", "name": {"given": "Test", "family": "User"}},
    f"{API}/v1/users/_9_1/courses": {"results": [
        {"courseRoleId": "Student", "course": {"id": "_1_1", "courseId": "CS-457-001-FA26", "name": "CS 457 Networks", "ultraStatus": "Ultra"}},
        {"courseRoleId": "Student", "course": {"id": "_2_1", "courseId": "MATH-301", "name": "Discrete Math"}},
    ]},
    f"{API}/v1/calendars/items": {"results": [
        {"type": "Course", "calendarId": "_1_1", "title": "Lecture", "location": "SEC 1001",
         "start": iso(NOW + timedelta(days=1)), "end": iso(NOW + timedelta(days=1, hours=1, minutes=15))},
    ]},
    f"{API}/v2/courses/_1_1/gradebook/columns": {"results": [
        {"id": "_c1", "name": "Lab 3: Wireshark", "contentId": "_a1", "score": {"possible": 50},
         "grading": {"due": iso(NOW + timedelta(days=2))}},
        {"id": "_c2", "name": "Quiz 2", "score": {"possible": 10}, "grading": {"due": iso(NOW - timedelta(days=1))}},
        {"id": "_c3", "name": "Total", "externalGrade": True, "score": {"possible": 100}},
    ]},
    f"{API}/v2/courses/_1_1/gradebook/users/_9_1": {"results": [
        {"columnId": "_c3", "displayGrade": {"score": 91.5}},
    ]},
    f"{API}/v2/courses/_1_1/gradebook/columns/_c1": {"id": "_c1", "name": "Lab 3: Wireshark", "contentId": "_a1",
        "score": {"possible": 50}, "grading": {"due": iso(NOW + timedelta(days=2))}},
    f"{API}/v1/courses/_1_1/contents/_a1": {"id": "_a1", "title": "Lab 3", "body": "<p>Capture a <b>TCP</b> handshake.</p>"},
    f"{API}/v1/courses/_1_1/contents/_a1/attachments": {"results": [{"id": "_att9", "fileName": "lab3.pdf"}]},
    f"{API}/v1/courses/_1_1/announcements": {"results": [
        {"title": "Exam moved", "body": "<p>Midterm is now Oct 15.</p>", "modified": iso(NOW - timedelta(days=1))}]},
    f"{API}/v1/courses/_1_1/contents": {"results": [
        {"id": "_f1", "title": "Week 5", "hasChildren": True, "contentHandler": {"id": "resource/x-bb-folder"}},
        {"id": "_a1", "title": "Lab 3", "contentHandler": {"id": "resource/x-bb-asmt-test-link"}},
    ]},
    f"{API}/v1/courses/_1_1/contents/_f1/children": {"results": [
        {"id": "_s1", "title": "Lecture 9 slides", "contentHandler": {"id": "resource/x-bb-file"}},
    ]},
    f"{API}/v1/courses/_1_1/contents/_s1/attachments": {"results": [{"id": "_att1", "fileName": "lec9-tcp.pptx"}]},
    f"{API}/v1/courses/_1_1/contents/_s1/attachments/_att1": {"id": "_att1", "fileName": "lec9-tcp.pptx"},
    f"{API}/v1/courses/_1_1/contents/_a1/attachments/_att9": {"id": "_att9", "fileName": "lab3.pdf"},
}
FILES = {
    f"{API}/v1/courses/_1_1/contents/_s1/attachments/_att1/download": PPTX,
    f"{API}/v1/courses/_1_1/contents/_a1/attachments/_att9/download": PDF,
}


def handler(req: httpx.Request) -> httpx.Response:
    assert "BbRouter" in req.headers.get("cookie", "")
    p = req.url.path
    if p in FILES:
        return httpx.Response(200, content=FILES[p], headers={"content-type": "application/octet-stream"})
    if p in ROUTES:
        return httpx.Response(200, json=ROUTES[p])
    if p.endswith("/attachments") or p.endswith("/announcements"):
        return httpx.Response(200, json={"results": []})
    return httpx.Response(404, json={"status": 404})


async def main():
    server._client = BlackboardClient("https://bb.test", cookie="BbRouter=x", transport=httpx.MockTransport(handler))
    s = server
    checks = [
        ("whoami", await s.whoami(), "Test User"),
        ("list_courses", await s.list_courses(), "CS 457 Networks"),
        ("schedule", await s.get_class_schedule(7), "SEC 1001"),
        ("upcoming", await s.get_upcoming(14, "cs457"), "Lab 3: Wireshark"),
        ("overdue", await s.get_upcoming(14, "_1_1"), "OVERDUE"),
        ("details", await s.get_assignment_details("CS 457", "_c1"), "Capture a TCP handshake"),
        ("grades", await s.get_grades("_1_1"), "91.5/100"),
        ("announce", await s.get_announcements("_1_1"), "Midterm is now Oct 15"),
        ("browse", await s.browse_course_content("_1_1"), "Lecture 9 slides"),
        ("find", await s.find_study_materials("_1_1", "week 5"), "lec9-tcp.pptx"),
        ("read pptx", await s.read_course_file("_1_1", "_s1"), "Remember AIMD"),
        ("read range", await s.read_course_file("_1_1", "_s1", start=2), "Slide 2: Fast Retransmit"),
        ("read pdf", await s.read_course_file("_1_1", "_a1", "_att9"), "scanned PDF"),
        ("search", await s.search_course_materials("_1_1", "duplicate ACKs"), "start=2"),
        ("ambiguous", await s.get_grades("nonexistent"), "No course matches"),
    ]
    ok = True
    for name, out, want in checks:
        good = want in out
        ok &= good
        print(f"{'PASS' if good else 'FAIL'} {name}")
        if not good:
            print(out)
    print()
    print(await s.read_course_file("_1_1", "_s1"))
    print()
    print(await s.get_upcoming(14))
    sys.exit(0 if ok else 1)


asyncio.run(main())
