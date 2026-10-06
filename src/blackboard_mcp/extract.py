"""Turn course files (slides, PDFs, Word docs) into study-ready text."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass


@dataclass
class Section:
    label: str  # e.g. "Slide 4" / "Page 12"
    title: str
    body: str
    notes: str = ""


def _clean(s: str) -> str:
    s = s.replace("\x0b", "\n").replace("\r", "")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def sniff_kind(data: bytes, filename: str = "", content_type: str = "") -> str:
    name = filename.lower()
    if data[:4] == b"%PDF" or name.endswith(".pdf") or "pdf" in content_type:
        return "pdf"
    if data[:2] == b"PK":
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            names = []
        if any(n.startswith("ppt/") for n in names):
            return "pptx"
        if any(n.startswith("word/") for n in names):
            return "docx"
    for ext in ("pptx", "docx", "txt", "md", "html", "htm"):
        if name.endswith("." + ext):
            return ext
    if content_type.startswith("text/"):
        return "txt"
    return "unknown"


def extract_pptx(data: bytes) -> list[Section]:
    from pptx import Presentation
    from pptx.util import Emu  # noqa: F401  (import check)

    prs = Presentation(io.BytesIO(data))
    out: list[Section] = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = _clean(slide.shapes.title.text_frame.text)
        lines: list[str] = []

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:  # group
                    walk(sh.shapes)
                    continue
                if slide.shapes.title is not None and sh.shape_id == slide.shapes.title.shape_id:
                    continue
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        t = "".join(r.text for r in p.runs).strip()
                        if t:
                            lines.append("  " * p.level + "- " + t)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        lines.append("| " + " | ".join(_clean(c.text) for c in row.cells) + " |")
                if sh.shape_type == 13:  # picture
                    alt = sh._element.xpath(".//p:cNvPr/@descr")
                    lines.append(f"[image{': ' + alt[0] if alt and alt[0] else ''}]")

        walk(slide.shapes)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
        out.append(Section(f"Slide {i}", title, _clean("\n".join(lines)), notes))
    return out


def extract_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = _clean(page.extract_text() or "")
        except Exception as e:  # malformed pages shouldn't kill the whole deck
            text = f"[could not extract text: {e}]"
        first = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
        out.append(Section(f"Page {i}", first[:120], text))
    return out


def extract_docx(data: bytes) -> list[Section]:
    import docx

    d = docx.Document(io.BytesIO(data))
    sections: list[Section] = []
    cur_title, buf = "", []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if p.style is not None and p.style.name.lower().startswith("heading"):
            if buf or cur_title:
                sections.append(Section(f"Section {len(sections) + 1}", cur_title, _clean("\n".join(buf))))
            cur_title, buf = t, []
        else:
            buf.append(t)
    for tbl in d.tables:
        for row in tbl.rows:
            buf.append("| " + " | ".join(_clean(c.text) for c in row.cells) + " |")
    sections.append(Section(f"Section {len(sections) + 1}", cur_title, _clean("\n".join(buf))))
    return sections


def extract_html(data: bytes) -> str:
    text = data.decode("utf-8", "replace")
    text = re.sub(r"(?is)<(script|style).*?</\1>", "", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>|</div>", "\n", text)
    text = re.sub(r"<[^>]+>", "", text)
    import html

    return _clean(html.unescape(text))


def extract(data: bytes, filename: str = "", content_type: str = "") -> tuple[str, list[Section]]:
    kind = sniff_kind(data, filename, content_type)
    if kind == "pptx":
        return kind, extract_pptx(data)
    if kind == "pdf":
        return kind, extract_pdf(data)
    if kind == "docx":
        return kind, extract_docx(data)
    if kind in ("html", "htm"):
        return kind, [Section("Document", filename, extract_html(data))]
    if kind in ("txt", "md"):
        return kind, [Section("Document", filename, _clean(data.decode("utf-8", "replace")))]
    if kind == "unknown" and filename.lower().endswith(".ppt"):
        raise ValueError("Legacy .ppt files aren't supported; ask for a .pptx or PDF export.")
    raise ValueError(f"Unsupported file type for '{filename}' ({content_type or 'unknown'}).")


def render(sections: list[Section], include_notes: bool = True) -> str:
    parts = []
    for s in sections:
        head = f"## {s.label}" + (f": {s.title}" if s.title else "")
        block = [head]
        if s.body:
            block.append(s.body)
        if include_notes and s.notes:
            block.append(f"_Speaker notes:_ {s.notes}")
        parts.append("\n".join(block))
    return "\n\n".join(parts)


def html_to_text(s: str | None) -> str:
    if not s:
        return ""
    return extract_html(s.encode())
